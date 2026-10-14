import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from physics_assistant.bank import connect, import_slides, records, match

class WorkflowTests(unittest.TestCase):
    def test_pptx_import_repeatability_and_shuffled_answers(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'lecture.pptx'
            slides = Presentation()
            slide = slides.slides.add_slide(slides.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text = 'What is the net force on a 2 kg object accelerating at 3 m/s²?\nA. 6 N\nB. 1.5 N\nC. 5 N'
            slides.save(path)
            with connect(Path(folder) / 'bank.db') as db:
                self.assertEqual(import_slides(db, path, Path(folder)/'images')['added'], 1)
                self.assertEqual(import_slides(db, path, Path(folder)/'images')['added'], 0)
                q = records(db)[0]
                db.execute('UPDATE questions SET answer=?,explanation=?,verified=1 WHERE id=?', ('6 N', 'F = ma = 2 × 3 = 6 N', q['id']))
                result = match(records(db), q['wording'], ['5 N','6 N','1.5 N'])
                self.assertEqual(result['answer'], '6 N')
                self.assertTrue(result['verified'])
                self.assertEqual(result['status'], 'exact')
                self.assertEqual(match(records(db), q['wording'], ['6 N','2 N'])['status'], 'manual_review')
                self.assertEqual(match(records(db), 'What is the velocity?', [])['status'], 'manual_review')
                self.assertEqual(match(records(db)*2, q['wording'], q['options'])['status'], 'manual_review')

    def test_pdf_extraction(self):
        # Minimal real PDF with text content, offsets calculated rather than hardcoded.
        with tempfile.TemporaryDirectory() as folder:
            objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
                       b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
                       b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
                       b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
            stream = b'BT /F1 12 Tf 50 700 Td (What is force?) Tj 0 -20 Td (A. Mass times acceleration) Tj 0 -20 Td (B. Speed) Tj ET'
            objects.append(b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream')
            content=b'%PDF-1.4\n';offsets=[0]
            for i,obj in enumerate(objects,1):
                offsets.append(len(content));content+=str(i).encode()+b' 0 obj\n'+obj+b'\nendobj\n'
            start=len(content);content+=b'xref\n0 6\n0000000000 65535 f \n'
            for offset in offsets[1:]:content+=f'{offset:010d} 00000 n \n'.encode()
            content+=f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n'.encode()
            path=Path(folder)/'lecture.pdf';path.write_bytes(content)
            with connect(Path(folder)/'bank.db') as db:
                self.assertEqual(import_slides(db,path,Path(folder)/'images')['added'],1)
                self.assertEqual(records(db)[0]['options'],['Mass times acceleration','Speed'])

    def test_fuzzy_and_unprepared(self):
        item={'wording':'What is the net force on this object?', 'options':['6 N','3 N'], 'answer':'', 'explanation':'', 'verified':0, 'source':'slides','page':1}
        result=match([item], 'What is the net force on this object ?', ['3 N','6 N'])
        self.assertEqual(result['status'],'uncertain_fuzzy')
        self.assertFalse(result['verified'])
        self.assertEqual(result['answer'],'')

if __name__ == '__main__':
    unittest.main()
