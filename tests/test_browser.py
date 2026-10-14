"""Full upload → review → same-tab simulator → read-only match integration."""
import os
import tempfile
import threading
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from playwright.sync_api import sync_playwright
from werkzeug.serving import make_server
from physics_assistant import web
from physics_assistant.bank import connect, records, match
from physics_assistant.browser import read_simulator

@unittest.skipUnless(os.environ.get('RUN_BROWSER_TESTS') == '1', 'Opt in with RUN_BROWSER_TESTS=1; requires Chromium')
class BrowserWorkflow(unittest.TestCase):
    def test_upload_review_navigation_and_read(self):
        with tempfile.TemporaryDirectory() as folder:
            previous = web.DATA
            web.DATA = Path(folder)
            server = make_server('127.0.0.1', 0, web.app)
            origin = f'http://127.0.0.1:{server.server_port}'
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                slides = Presentation()
                slide = slides.slides.add_slide(slides.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(4)).text = 'What is force?\nA. Mass times acceleration\nB. Speed'
                file = Path(folder)/'demo.pptx';slides.save(file)
                with sync_playwright() as p:
                    browser = p.chromium.launch(headless=True, executable_path=os.environ.get('CHROMIUM_EXECUTABLE'))
                    page = browser.new_page()
                    page.goto(origin)
                    page.locator('input[type=file]').set_input_files(file)
                    page.get_by_role('button', name='Import slides').click()
                    page.locator('article').wait_for()
                    page.locator('select').select_option('Mass times acceleration')
                    page.locator('textarea').fill('Newton’s second law: F = ma.')
                    page.get_by_label('Verified by me').check()
                    with page.expect_response(lambda r: '/api/questions/' in r.url and r.request.method == 'POST'):
                        page.get_by_role('button', name='Save answer').click()
                    page.goto(origin+'/simulator')
                    page.get_by_role('link',name='What is force?').click()
                    page.locator('[data-question]').wait_for()
                    current=read_simulator(page, origin)
                    self.assertEqual(current['choices'],['Speed','Mass times acceleration'])
                    with connect(Path(folder)/'bank.sqlite3') as db:
                        result=match(records(db),current['wording'],current['choices'])
                    self.assertEqual(result['answer'],'Mass times acceleration')
                    self.assertTrue(result['verified'])
                    self.assertEqual(page.locator('input:checked').count(),0)
                    page.get_by_role('link',name='Back to available questions').click()
                    self.assertIsNone(read_simulator(page, origin))
                    self.assertEqual(page.request.post(origin+'/api/match', data={}).status,403)
                    browser.close()
            finally:
                server.shutdown()
                thread.join(timeout=5)
                web.DATA=previous
