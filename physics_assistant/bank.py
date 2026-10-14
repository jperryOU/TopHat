"""Local slide extraction and conservative, answer-text matching."""
import hashlib
import json
import re
import sqlite3
import unicodedata
from pathlib import Path
from rapidfuzz.fuzz import ratio


def normalize(text):
    return ' '.join(unicodedata.normalize('NFKC', text).casefold().split())


def connect(path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(path)
    db.row_factory = sqlite3.Row
    db.execute('''CREATE TABLE IF NOT EXISTS questions (
        id TEXT PRIMARY KEY, source TEXT, page INTEGER, wording TEXT,
        options TEXT, context TEXT, images TEXT, answer TEXT DEFAULT '',
        explanation TEXT DEFAULT '', verified INTEGER DEFAULT 0)''')
    return db


def records(db):
    result = []
    for row in db.execute('SELECT * FROM questions ORDER BY source,page,id'):
        item = dict(row)
        item['options'] = json.loads(item['options'])
        item['images'] = json.loads(item['images'])
        result.append(item)
    return result


OPTION = re.compile(r'^\s*(?:\(?([A-Ha-h])\s*[).:]|([1-8])\s*[).])\s+(.+)$')


def split_questions(text):
    """Conservative heuristic: question marks or explicit Question/Q labels.

    Preserve the extracted wording; retain the whole slide separately as context.
    Text-only statements without a question marker need manual review.
    """
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    results, prompt, options = [], [], []
    def flush():
        if prompt and ('?' in ' '.join(prompt) or re.match(r'^(question\b|q\d+\b)', prompt[0], re.I)):
            results.append(('\n'.join(prompt), options.copy()))
    for line in lines:
        option = OPTION.match(line)
        if option and prompt:
            options.append(option.group(3))
        elif options:
            if '?' in line or re.match(r'^(question\b|q\d+\b)', line, re.I):
                flush()
                prompt, options = [line], []
            else:
                options[-1] += '\n' + line
        else:
            if prompt and '?' in ' '.join(prompt):
                flush()
                prompt = []
            prompt.append(line)
    flush()
    return results


def extract(path, image_dir):
    path, image_dir = Path(path), Path(image_dir)
    image_dir.mkdir(parents=True, exist_ok=True)
    prefix = hashlib.sha256(path.read_bytes()).hexdigest()[:16]
    pages = []
    if path.suffix.lower() == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        for number, page in enumerate(reader.pages, 1):
            images = []
            for index, img in enumerate(page.images):
                name = f'{prefix}-{number}-{index}{Path(img.name).suffix}'
                (image_dir / name).write_bytes(img.data)
                images.append(name)
            pages.append((number, page.extract_text() or '', images))
    elif path.suffix.lower() == '.pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        def walk(shapes):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from walk(shape.shapes)
                else:
                    yield shape
        for number, slide in enumerate(Presentation(path).slides, 1):
            texts, images = [], []
            for index, shape in enumerate(walk(slide.shapes)):
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    name = f'{prefix}-{number}-{index}.{shape.image.ext}'
                    (image_dir / name).write_bytes(shape.image.blob)
                    images.append(name)
            pages.append((number, '\n'.join(texts), images))
    else:
        raise ValueError('Use PDF or PPTX. Convert older .ppt files to .pptx first.')
    return pages


def import_slides(db, path, image_dir, source=None):
    pages = extract(path, image_dir)
    count = 0
    for page, context, images in pages:
        for wording, options in split_questions(context):
            identity = hashlib.sha256(json.dumps([source or Path(path).name, page, wording, options]).encode()).hexdigest()
            cursor = db.execute('''INSERT OR IGNORE INTO questions
                (id,source,page,wording,options,context,images) VALUES (?,?,?,?,?,?,?)''',
                (identity, source or Path(path).name, page, wording, json.dumps(options), context, json.dumps(images)))
            count += cursor.rowcount
    db.commit()
    return {'added': count, 'pages': len(pages), 'pages_without_questions':
            [number for number, text, _ in pages if not split_questions(text)]}


def match(items, wording, choices):
    if not normalize(wording):
        return {'status': 'manual_review', 'reason': 'Empty question.'}
    ranked = sorted([(ratio(normalize(wording), normalize(q['wording'])), q) for q in items],
                    key=lambda pair: pair[0], reverse=True)
    if not ranked or ranked[0][0] < 94:
        return {'status': 'manual_review', 'reason': 'No reliable wording match.'}
    score, question = ranked[0]
    if len(ranked) > 1 and score - ranked[1][0] < 5:
        return {'status': 'manual_review', 'reason': 'Multiple similar questions; inspect the bank.'}
    # Requiring an identical choice set guards against negated or changed options.
    expected = sorted(normalize(c) for c in question['options'])
    actual = sorted(normalize(c) for c in choices)
    if expected != actual:
        return {'status': 'manual_review', 'reason': 'Answer choices differ from the slide.'}
    answer = question['answer']
    resolved = [c for c in choices if normalize(c) == normalize(answer)]
    if choices and len(resolved) != 1:
        answer = ''
    elif resolved:
        answer = resolved[0]
    return {'status': 'exact' if score == 100 else 'uncertain_fuzzy',
            'similarity': round(score, 1), 'verified': bool(question['verified']),
            'answer': answer, 'explanation': question['explanation'],
            'source': question['source'], 'page': question['page'],
            'reason': 'Review before selecting. Similarity is not a probability of correctness.'}
